#!/usr/bin/env python
from docx import Document
import os
import sys

# Fix encoding for Windows
sys.stdout.reconfigure(encoding='utf-8')

os.chdir(r'c:\Users\M.Gowthami\Downloads\Fusion_v3\23-System-Admin-specs')

def extract_doc_content(filename, title, output_file):
    output_file.write("\n" + "=" * 80 + "\n")
    output_file.write(f"EXTRACTING: {title}\n")
    output_file.write("=" * 80 + "\n")
    doc = Document(filename)
    
    # Extract paragraphs
    output_file.write("\n--- PARAGRAPHS ---\n")
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            output_file.write(text + "\n")
    
    # Extract tables
    output_file.write("\n--- TABLES ---\n")
    for i, table in enumerate(doc.tables):
        output_file.write(f"\nTable {i+1}:\n")
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            output_file.write(" | ".join(row_data) + "\n")

# Extract all documents to file
with open('extracted_specs.txt', 'w', encoding='utf-8') as f:
    extract_doc_content('super_admin_UC.docx', 'USE CASES', f)
    extract_doc_content('super_admin_wf.docx', 'WORKFLOWS', f)
    extract_doc_content('super_admin_BR.docx', 'BUSINESS RULES', f)

print("Extraction complete. Check extracted_specs.txt")
